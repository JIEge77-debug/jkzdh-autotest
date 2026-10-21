import openpyxl

from config.config import EXCEL_FILE, SHEET_NAME


def read_excel(file_path = EXCEL_FILE,sheet_name = SHEET_NAME):

    #打开Excel文件
    workbook = openpyxl.load_workbook(file_path) #参数文件路径

    #选择表
    worksheet = workbook[sheet_name] #表名严格区分大小写，文件里的表名是 Sheet1
    # 也可以不写死表名，直接取第一个表：
    # worksheet = workbook[workbook.sheetnames[0]]
    #读数据操作 zip() 函数的作用是将多个可迭代对象中的元素按顺序组合成一个元组
    data = [] #空列表，用于组装字典
    key = [cell.value for cell in worksheet[2]]  #拿key行，拿第二行，生成一个key列表
    for row in worksheet.iter_rows(min_row=3,values_only=True): #从第三行开始拿，只返回值
        dict_data = dict(zip(key, row))
        #如果读取的 is_true 字段的值是TRUE，则append，否则不append
        # print(dict_data["is_True"])
        if dict_data["is_true"]:
            data.append(dict_data)
        # data.append(dict_data)
    #关闭Excel文件
    workbook.close()
    return data

read_excel()
